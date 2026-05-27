import re
import sys
from pathlib import Path

try:
    import tkinter as tk
    from tkinter import filedialog, messagebox, scrolledtext
    GUI_AVAILABLE = True
except ImportError:
    tk = None
    filedialog = None
    messagebox = None
    scrolledtext = None
    GUI_AVAILABLE = False

from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException

FIELD_PATTERNS = [
    ("Company Name", r"(?:Company Name|Company)\s*[:：]?\s*(.+)"),
    ("Product Info", r"(?:Product Info|Product)\s*[:：]?\s*(.+)"),
    ("Purchase No.", r"(?:Purchase No\.?|Purchase Number|Purchase)\s*[:：]?\s*(.+)"),
    ("Invoice No.", r"(?:Invoice No\.?|Invoice Number|Invoice)\s*[:：]?\s*(.+)"),
    ("Delivery Date", r"(?:Delivery Date|Delivery)\s*[:：]?\s*(.+)"),
    ("Vehicle No.", r"(?:Vehicle No\.?|Vehicle Number|Vehicle)\s*[:：]?\s*(.+)"),
    ("Phone", r"(?:Phone|Phone No\.?|Telephone)\s*[:：]?\s*(.+)"),
    ("Name", r"(?:Name|Driver Name|Contact Name)\s*[:：]?\s*(.+)"),
    ("ID Card No.", r"(?:ID Card No\.?|ID Card Number|ID No\.?)\s*[:：]?\s*(.+)"),
    ("Tonnage", r"(?:Tonnage|Weight)\s*[:：]?\s*(.+)"),
]

HEADER_ROW = [field for field, _ in FIELD_PATTERNS]


def extract_fields(message_text: str) -> dict:
    message_text = message_text.strip()
    result = {field: "" for field, _ in FIELD_PATTERNS}
    if not message_text:
        return result

    lines = [line.strip() for line in message_text.splitlines() if line.strip()]
    normalized_text = message_text.replace("：", ":")

    for field, pattern in FIELD_PATTERNS:
        regex = re.compile(pattern, re.IGNORECASE)
        for line in lines:
            match = regex.search(line)
            if match:
                result[field] = match.group(1).strip()
                break
        else:
            match = regex.search(normalized_text)
            if match:
                result[field] = match.group(1).strip()

    return result


def open_or_create_workbook(excel_path: Path, sheet_name: str | None = None):
    if excel_path.exists():
        try:
            workbook = load_workbook(excel_path)
        except InvalidFileException as exc:
            raise ValueError("Selected file is not a valid .xlsx workbook.") from exc
    else:
        workbook = Workbook()

    if sheet_name:
        if sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
        else:
            sheet = workbook.create_sheet(title=sheet_name)
    else:
        sheet = workbook.active

    if sheet.max_row == 1 and all(cell.value is None for cell in sheet[1]):
        sheet.append(HEADER_ROW)

    return workbook, sheet


def append_to_excel(excel_file: str, row_values: list[str], sheet_name: str | None = None) -> None:
    path = Path(excel_file)
    workbook, sheet = open_or_create_workbook(path, sheet_name)
    sheet.append(row_values)
    workbook.save(path)


def format_extracted_values(values: dict) -> str:
    return "\n".join(f"{key}: {value}" for key, value in values.items())


def choose_file(file_var) -> None:
    file_path = filedialog.asksaveasfilename(
        title="Choose or create Excel file",
        defaultextension=".xlsx",
        filetypes=[("Excel Files", "*.xlsx")],
        initialfile="delivery.xlsx",
    )
    if file_path:
        file_var.set(file_path)


def clear_form(message_box, output_box, status_label, sheet_var) -> None:
    message_box.delete("1.0", tk.END)
    output_box.configure(state="normal")
    output_box.delete("1.0", tk.END)
    output_box.configure(state="disabled")
    sheet_var.set("")
    status_label.config(text="Ready.", fg="green")


def append_action(
    message_box,
    excel_path_var,
    sheet_name_var,
    output_box,
    status_label,
) -> None:
    message_text = message_box.get("1.0", tk.END).strip()
    excel_file = excel_path_var.get().strip()
    sheet_name = sheet_name_var.get().strip() or None

    if not message_text:
        status_label.config(text="Please paste a delivery message before appending.", fg="red")
        return

    if not excel_file:
        status_label.config(text="Please choose an Excel file.", fg="red")
        return

    if not excel_file.lower().endswith(".xlsx"):
        status_label.config(text="Please choose a .xlsx workbook.", fg="red")
        return

    try:
        extracted = extract_fields(message_text)
        values_row = [extracted[field] for field, _ in FIELD_PATTERNS]
        append_to_excel(excel_file, values_row, sheet_name)

        output_box.configure(state="normal")
        output_box.delete("1.0", tk.END)
        output_box.insert(tk.END, format_extracted_values(extracted))
        output_box.configure(state="disabled")

        status_label.config(text="Data appended successfully.", fg="green")
    except Exception as exc:
        status_label.config(text=f"Error: {exc}", fg="red")
        messagebox.showerror("Append Error", str(exc))


def cli_main() -> None:
    print("Tkinter is not available in this Python environment.")
    print("This tool can still run in command-line mode, or use a Python installation with Tk support.")
    print("\nPaste the delivery message below. Enter a blank line to finish:")

    message_lines = []
    while True:
        try:
            line = input()
        except EOFError:
            break
        if line.strip() == "" and message_lines:
            break
        message_lines.append(line)

    message_text = "\n".join(message_lines).strip()
    if not message_text:
        print("No message provided. Exiting.")
        return

    excel_file = input("Excel file path (.xlsx): ").strip()
    if not excel_file:
        print("No Excel file selected. Exiting.")
        return

    sheet_name = input("Sheet name (optional): ").strip() or None

    try:
        extracted = extract_fields(message_text)
        values_row = [extracted[field] for field, _ in FIELD_PATTERNS]
        append_to_excel(excel_file, values_row, sheet_name)
        print("Data appended successfully.")
        print("\nExtracted values:")
        print(format_extracted_values(extracted))
    except Exception as exc:
        print(f"Error: {exc}")


def build_window() -> None:
    root = tk.Tk()
    root.title("Delivery Message to Excel")
    root.geometry("940x720")
    root.minsize(800, 640)

    excel_path_var = tk.StringVar()
    sheet_name_var = tk.StringVar()

    top_frame = tk.Frame(root, padx=12, pady=12)
    top_frame.pack(fill=tk.BOTH, expand=True)

    tk.Label(top_frame, text="Paste one delivery message below:", font=(None, 11)).pack(anchor="w")
    message_box = tk.Text(top_frame, wrap="word", height=12)
    message_box.pack(fill=tk.BOTH, expand=True, pady=(4, 10))

    file_frame = tk.Frame(top_frame)
    file_frame.pack(fill=tk.X, pady=(0, 8))
    tk.Label(file_frame, text="Excel file:").pack(side=tk.LEFT)
    tk.Entry(file_frame, textvariable=excel_path_var, width=70).pack(side=tk.LEFT, padx=(8, 8), fill=tk.X, expand=True)
    tk.Button(file_frame, text="Choose file", command=lambda: choose_file(excel_path_var)).pack(side=tk.LEFT)

    sheet_frame = tk.Frame(top_frame)
    sheet_frame.pack(fill=tk.X, pady=(0, 10))
    tk.Label(sheet_frame, text="Sheet name (optional, default first tab):").pack(side=tk.LEFT)
    tk.Entry(sheet_frame, textvariable=sheet_name_var, width=30).pack(side=tk.LEFT, padx=(8, 0))

    action_frame = tk.Frame(top_frame)
    action_frame.pack(fill=tk.X, pady=(0, 10))
    tk.Button(action_frame, text="Append to Excel", width=16, command=lambda: append_action(message_box, excel_path_var, sheet_name_var, output_box, status_label)).pack(side=tk.LEFT)
    tk.Button(action_frame, text="Clear", width=12, command=lambda: clear_form(message_box, output_box, status_label, sheet_name_var)).pack(side=tk.LEFT, padx=(10, 0))
    tk.Button(action_frame, text="Exit", width=12, command=root.destroy).pack(side=tk.LEFT, padx=(10, 0))

    status_frame = tk.Frame(top_frame)
    status_frame.pack(fill=tk.X, pady=(0, 4))
    tk.Label(status_frame, text="Status:", font=(None, 10)).pack(side=tk.LEFT)
    status_label = tk.Label(status_frame, text="Ready.", fg="green")
    status_label.pack(side=tk.LEFT, padx=(8, 0))

    tk.Label(top_frame, text="Extracted values:").pack(anchor="w")
    output_box = scrolledtext.ScrolledText(top_frame, wrap="word", height=14, state="disabled")
    output_box.pack(fill=tk.BOTH, expand=True)

    root.mainloop()


def main() -> None:
    if GUI_AVAILABLE:
        build_window()
    else:
        cli_main()


if __name__ == "__main__":
    main()
